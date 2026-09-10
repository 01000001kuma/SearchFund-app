"""Tests de generación finance-grade de PDF y Excel (engines.export)."""
from io import BytesIO

from backend.engines.export import export_pdf, export_excel, score_band
from tests.conftest import make_company


def _pdf_header(buf: bytes) -> bool:
    """Un PDF válido empieza por %PDF."""
    return buf.startswith(b"%PDF")


def _sample_company(cif="B85767044", name="ALIMENTACION ANIMAL NANTA SL"):
    c = make_company(
        cif=cif, name=name, slug=name.lower().replace(" ", "-"),
        website="https://nanta.es", phone="+34 918075400",
        email="nanta@nutreco.com", ebitda=2_000_000, revenue=12_000_000,
    )
    c.score = 55.0
    c.score_breakdown = {
        "total": 55, "borme": 55, "financial": None,
        "interpretation": "CANDIDATO MODERADO",
        "breakdown": {"borme": {
            "admin_age": 0, "stability": 0, "family": 20,
            "no_council": 10, "cnae": 0, "recent_activity": 25,
        }},
    }
    return c


def test_score_band_alineado_con_umbrales():
    assert score_band(93)[1] == "MUY BUENO"
    assert score_band(65)[1] == "BUENO"
    assert score_band(55)[1] == "MODERADO"
    assert score_band(25)[1] == "BAJO"
    assert score_band(5)[1] == "NO RECOMENDADO"
    assert score_band(None)[1] == "Sin evaluar"


def test_export_pdf_genera_pdf_valido():
    pdf = export_pdf(_sample_company())
    assert isinstance(pdf, bytes)
    assert len(pdf) > 2000
    assert pdf.startswith(b"%PDF")


def test_export_pdf_sin_datos_extra_no_peta():
    company = make_company(ebitda=None, revenue=None, website=None, phone=None, email=None)
    pdf = export_pdf(company)
    assert pdf.startswith(b"%PDF")


def test_export_excel_genera_xlsx_validos():
    companies = [
        _sample_company(),
        make_company(cif="A46103834", name="MERCADONA SA", slug="mercadona", ebitda=None, revenue=None),
    ]
    xlsx = export_excel(companies)
    assert isinstance(xlsx, bytes)
    assert xlsx.startswith(b"PK")
    assert len(xlsx) > 1000


def test_export_excel_estructura_finance_grade():
    from openpyxl import load_workbook

    companies = [
        _sample_company(),
        make_company(cif="A46103834", name="MERCADONA SA", slug="mercadona", ebitda=None, revenue=None),
    ]
    wb = load_workbook(BytesIO(export_excel(companies)))
    assert "Candidatas" in wb.sheetnames
    assert "Leyenda" in wb.sheetnames
    ws = wb["Candidatas"]
    # fila 1 título, 2 generación, 3 vacía, 4 cabecera, 5-6 datos
    assert ws.max_row == 6
    headers = [c.value for c in ws[4]]
    for key in ("CIF", "Nombre", "EBITDA (€)", "Facturación (€)", "Score", "Clasificación"):
        assert key in headers
    # cifras numéricas, no strings formateados
    assert ws.cell(row=5, column=11).value == 2_000_000
    assert ws.cell(row=6, column=11).value is None
    # score con formato de banda
    assert ws.cell(row=5, column=17).value == "MODERADO"
    assert ws.cell(row=5, column=16).fill.start_color.rgb is not None
    # autofiltro y panes congelados
    assert ws.auto_filter.ref is not None
    assert ws.freeze_panes == "C5"
    # hoja Leyenda con bandas
    lg = wb["Leyenda"]
    assert lg["A4"].value == "MUY BUENO"
    assert lg["B4"].value == "80–100"


def test_export_excel_cabecera_coincide_con_datos():
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(export_excel([_sample_company()])))
    ws = wb["Candidatas"]
    headers = [c.value for c in ws[4]]
    row = [c.value for c in ws[5]]
    assert len(headers) == len(row)


def test_export_pdf_sin_reprs_de_objetos():
    """Los flowables deben dibujarse, no imprimirse como repr() (regresión)."""
    pdf = export_pdf(_sample_company())
    assert b"ScoreGauge" not in pdf
    assert b"_Bar object" not in pdf
    assert b"backend.engines" not in pdf


def test_export_pdf_calcula_desglose_si_no_esta_persistido():
    from backend.models.financial import FinancialData
    c = make_company(cif="B22222222", name="SIN DESGLOSE SL", slug="sin-desglose")
    c.score = 45.0
    c.score_breakdown = None  # sin desglose persistido
    pdf = export_pdf(c)
    assert pdf.startswith(b"%PDF")
